#!/usr/bin/env python3
"""Gera/atualiza a planilha de textos das artes conceituais (docs/concept_art_texts.xlsx).

- Uma linha por arquivo de assets/concepts (categoria deduzida do nome).
- Colunas editáveis: titulo_pt, texto_pt (EN será traduzido depois: titulo_en, texto_en ficam vazias).
- Reexecutar é seguro: linhas já preenchidas são preservadas; só entram arquivos novos.

Uso: ./venv/bin/python tools/generate_concept_texts_sheet.py
Requer openpyxl (apenas ferramenta de desenvolvimento; não é dependência do jogo).
"""
import os
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT)

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

CONCEPTS_DIR = os.path.join(ROOT, "assets", "concepts")
OUT_PATH = os.path.join(ROOT, "docs", "concept_art_texts.xlsx")
HEADERS = ["arquivo", "categoria", "nome_sugerido", "titulo_pt", "texto_pt", "titulo_en", "texto_en", "status"]
STATUSES = ["pendente", "rascunho", "revisado", "traduzido"]
IMG_EXT = (".jpg", ".jpeg", ".png")


def classify(filename: str) -> str:
    """Categoria deduzida do nome do arquivo."""
    n = filename.lower()
    if n.endswith("_full_concept.jpg"):
        return "personagem"
    if "_duel" in n:
        return "duelo"
    if "logo" in n or "ending" in n or "silhouette" in n:
        return "logo/final"
    if n.startswith(("gashadokuro",)):
        return "chefe"
    if n.endswith("_concept.jpg"):
        return "arena"
    return "outros"


def suggest_name(filename: str) -> str:
    stem = os.path.splitext(filename)[0]
    for suf in ("_full_concept", "_concept"):
        if stem.endswith(suf):
            stem = stem[: -len(suf)]
    return stem.replace("_", " ").title()


def list_concepts() -> list[str]:
    files = [f for f in os.listdir(CONCEPTS_DIR) if f.lower().endswith(IMG_EXT)]
    order = {"personagem": 0, "arena": 1, "duelo": 2, "chefe": 3, "logo/final": 4, "outros": 5}
    return sorted(files, key=lambda f: (order[classify(f)], f.lower()))


def read_existing() -> dict[str, list]:
    if not os.path.isfile(OUT_PATH):
        return {}
    ws = load_workbook(OUT_PATH).active
    rows = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            rows[str(row[0])] = list(row) + [None] * (len(HEADERS) - len(row))
    return rows


def build() -> int:
    existing = read_existing()
    wb = Workbook()
    ws = wb.active
    ws.title = "artes_conceituais"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3B2F2F")
        c.alignment = Alignment(vertical="center")
    for f in list_concepts():
        old = existing.get(f)
        if old:
            ws.append(old[: len(HEADERS)])
        else:
            ws.append([f, classify(f), suggest_name(f), "", "", "", "", "pendente"])
    for col, width in zip("ABCDEFGH", (34, 13, 26, 30, 70, 30, 70, 12)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"H2:H{ws.max_row}")
    ws.freeze_panes = "C2"
    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    wb.save(OUT_PATH)
    return ws.max_row - 1


if __name__ == "__main__":
    print(f"{build()} artes na planilha: {OUT_PATH}")
